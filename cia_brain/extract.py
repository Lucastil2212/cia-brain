from __future__ import annotations

import gzip
import html
import io
import json
import logging
import mimetypes
import os
import re
import subprocess
import tempfile
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
import trafilatura
from bs4 import BeautifulSoup
from docx import Document
from pptx import Presentation
from pypdf import PdfReader

from .models import ExtractedEvent
from .settings import Settings

log = logging.getLogger(__name__)
SPACE_RE = re.compile(r"[ \t]+")


def utcnow() -> str:
    return datetime.now(timezone.utc).isoformat()


def clean_text(text: str) -> str:
    text = html.unescape(text or "").replace("\x00", " ")
    lines = [SPACE_RE.sub(" ", line).strip() for line in text.splitlines()]
    return "\n".join(line for line in lines if line).strip()


def extract_html(path: Path) -> tuple[str, str]:
    raw = path.read_bytes()
    decoded = raw.decode("utf-8", errors="replace")
    soup = BeautifulSoup(decoded, "html.parser")
    title = clean_text(soup.title.get_text(" ") if soup.title else "")
    text = trafilatura.extract(decoded, include_comments=False, include_tables=True, favor_recall=True)
    if not text:
        for node in soup(["script", "style", "noscript"]):
            node.decompose()
        text = soup.get_text("\n")
    return title, clean_text(text)


def extract_pdf(path: Path, settings: Settings) -> tuple[str, str]:
    reader = PdfReader(str(path))
    title = ""
    if reader.metadata and getattr(reader.metadata, "title", None):
        title = str(reader.metadata.title)
    pages = [clean_text(page.extract_text() or "") for page in reader.pages]
    text = "\n\n".join(p for p in pages if p)
    if settings.ocr_enabled and len(text) < max(200, len(reader.pages) * 40):
        ocr = ocr_pdf(path, settings.ocr_max_pages)
        if len(ocr) > len(text):
            text = ocr
    return clean_text(title), clean_text(text)


def ocr_pdf(path: Path, max_pages: int) -> str:
    if not shutil_which("pdftoppm") or not shutil_which("tesseract"):
        return ""
    texts: list[str] = []
    with tempfile.TemporaryDirectory(prefix="cia-ocr-") as td:
        prefix = str(Path(td) / "page")
        cmd = ["pdftoppm", "-jpeg", "-r", "170"]
        if max_pages > 0:
            cmd += ["-f", "1", "-l", str(max_pages)]
        cmd += [str(path), prefix]
        subprocess.run(cmd, check=False, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        for image in sorted(Path(td).glob("page-*.jpg")):
            proc = subprocess.run(
                ["tesseract", str(image), "stdout", "-l", "eng"],
                check=False,
                stdout=subprocess.PIPE,
                stderr=subprocess.DEVNULL,
            )
            texts.append(proc.stdout.decode("utf-8", errors="replace"))
    return clean_text("\n\n".join(texts))


def shutil_which(name: str) -> str | None:
    import shutil
    return shutil.which(name)


def extract_docx(path: Path) -> tuple[str, str]:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    title = doc.core_properties.title or ""
    return clean_text(title), clean_text("\n".join(parts))


def extract_pptx(path: Path) -> tuple[str, str]:
    prs = Presentation(str(path))
    parts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                parts.append(shape.text)
    return "", clean_text("\n".join(parts))


def extract_xlsx(path: Path) -> tuple[str, str]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            vals = [str(v) for v in row if v is not None]
            if vals:
                parts.append(" | ".join(vals))
    return "", clean_text("\n".join(parts))


def extract_zip(path: Path) -> tuple[str, str]:
    parts = []
    try:
        with zipfile.ZipFile(path) as zf:
            infos = zf.infolist()
            parts.append("Archive members:")
            for info in infos[:10000]:
                parts.append(f"{info.filename}\t{info.file_size}")
    except Exception:
        pass
    return "", clean_text("\n".join(parts))


def extract_text(path: Path, mime: str, settings: Settings) -> tuple[str, str]:
    if mime == "application/vnd.cia-brain.record+json":
        record = json.loads(path.read_text(encoding="utf-8"))
        return clean_text(record.get("title", "")), clean_text(record.get("text", ""))
    if mime.startswith("image/") and settings.ocr_enabled and shutil_which("tesseract"):
        proc = subprocess.run(["tesseract", str(path), "stdout", "-l", "eng"],
                              capture_output=True, timeout=120, check=True)
        return "", clean_text(proc.stdout.decode("utf-8", errors="replace"))
    if mime in {"application/json", "application/geo+json"}:
        return "", clean_text(json.dumps(json.loads(path.read_text()), ensure_ascii=False))
    suffix = path.suffix.lower()
    if mime in {"text/html", "application/xhtml+xml"} or suffix in {".html", ".htm"}:
        return extract_html(path)
    if mime == "application/pdf" or suffix == ".pdf":
        return extract_pdf(path, settings)
    if suffix == ".docx" or "wordprocessingml" in mime:
        return extract_docx(path)
    if suffix == ".pptx" or "presentationml" in mime:
        return extract_pptx(path)
    if suffix == ".xlsx" or "spreadsheetml" in mime:
        return extract_xlsx(path)
    if suffix == ".zip" or mime in {"application/zip", "application/x-zip-compressed"}:
        return extract_zip(path)
    if mime.startswith("text/") or suffix in {".txt", ".csv", ".xml", ".md", ".rtf"}:
        return "", clean_text(path.read_text(encoding="utf-8", errors="replace"))
    return "", ""


def normalize_document(payload: dict, settings: Settings) -> ExtractedEvent:
    from .paths import normalized_cas_path, path_under, require_sha256

    digest = require_sha256(payload.get("sha256", ""))
    src = path_under(payload["path"], settings.raw_dir)
    title, text = extract_text(src, payload.get("mime", ""), settings)
    source_urls = [payload["final_url"]]
    manifest_path = settings.manifests_dir / digest[:2] / f"{digest}.json"
    if manifest_path.exists():
        try:
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
            source_urls = sorted({
                x.get("final_url") or x.get("url")
                for x in manifest.get("sources", [])
                if x.get("final_url") or x.get("url")
            }) or source_urls
        except Exception:
            pass
    doc = {
        "source_url": payload["final_url"],
        "source_urls": source_urls,
        "requested_url": payload["url"],
        "source_sha256": digest,
        "mime": payload.get("mime", ""),
        "size": payload.get("size", 0),
        "fetched_at": payload.get("fetched_at"),
        "title": title,
        "text": text,
        "text_chars": len(text),
        "extracted_at": utcnow(),
    }
    if payload.get("mime") == "application/vnd.cia-brain.record+json":
        record = json.loads(src.read_text(encoding="utf-8"))
        for key in ("source_id", "agency", "modalities", "record_id", "published_at",
                    "geometry", "media", "properties", "measurements"):
            if key in record:
                doc[key] = record[key]
        doc["snapshot_sha256"] = payload.get("snapshot_sha256")
    for key in ("source_id", "agency", "modalities", "record_id"):
        if key in payload:
            doc[key] = payload[key]
    if payload.get("mime", "").startswith("image/") and not text:
        doc["title"] = title = payload.get("title", "Public source image")
        doc["text"] = text = title + " — image archived; visual content not interpreted."
        doc["text_chars"] = len(text)
    out = normalized_cas_path(settings, digest)
    out.parent.mkdir(parents=True, exist_ok=True)
    tmp = out.with_suffix(".tmp")
    with gzip.open(tmp, "wt", encoding="utf-8") as f:
        json.dump(doc, f, ensure_ascii=False)
    os.replace(tmp, out)
    if settings.graph_enabled:
        from .graph import KnowledgeGraph
        KnowledgeGraph(settings).ingest(doc)
    return ExtractedEvent(
        source_url=doc["source_url"],
        source_sha256=doc["source_sha256"],
        normalized_path=str(out),
        title=title,
        mime=doc["mime"],
        text_chars=len(text),
        extracted_at=doc["extracted_at"],
    )
